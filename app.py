"""상장회사 분석 웹 앱 (Streamlit)"""
import datetime as dt
import html
import io
import json
import os
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st

from concurrent.futures import ThreadPoolExecutor

import dart
import prices

# ※ 속도를 위해 dart.py·prices.py 는 앱을 켤 때 한 번만 읽음 (코드를 고치면 run.bat 을 다시 실행)

st.set_page_config(page_title="상장회사 분석", layout="wide")
# 작업 중 표시(달리는 아이콘)를 화면 가운데에 크게 표시
st.markdown("""<style>
[data-testid="stStatusWidget"] {
    position: fixed !important; top: 50% !important; left: 50% !important;
    transform: translate(-50%, -50%) scale(2.5); z-index: 999999;
    background: rgba(255,255,255,0.9); border-radius: 12px; padding: 6px 10px;
    box-shadow: 0 2px 12px rgba(0,0,0,0.15);
}
/* 좌우 스크롤바 크게 */
::-webkit-scrollbar { height: 16px; width: 12px; }
::-webkit-scrollbar-track { background: rgba(128,128,128,0.12); border-radius: 8px; }
::-webkit-scrollbar-thumb { background: rgba(128,128,128,0.55); border-radius: 8px; border: 3px solid transparent; background-clip: padding-box; }
::-webkit-scrollbar-thumb:hover { background: rgba(100,100,100,0.8); background-clip: padding-box; }
* { scrollbar-width: auto; }
/* 상단 공백 줄이기 */
.block-container, [data-testid="stMainBlockContainer"] { padding-top: 1rem !important; padding-bottom: 1rem !important; }
[data-testid="stHeader"] { height: 2rem !important; min-height: 2rem !important; background: transparent !important; }
[data-testid="stSidebarHeader"] { padding-top: 0.5rem !important; padding-bottom: 0 !important; height: auto !important; }
h3:first-of-type { margin-top: 0 !important; padding-top: 0 !important; }
</style>""", unsafe_allow_html=True)
ROOT = Path(__file__).parent
EOK = 1e8  # 억원
ALL = "전체"


def secret(name):
    try:
        v = st.secrets.get(name, "")
    except Exception:
        v = ""
    return v or os.environ.get(name, "")


# 미리 받아둔 자료 폴더: 앱 폴더의 cache, 없으면 한 단계 위(저장소 맨 위)의 cache
CACHE = next((p for p in (ROOT / "cache", ROOT.parent / "cache") if p.exists()), ROOT / "cache")


@st.cache_resource(show_spinner="미리 받아둔 자료 불러오는 중...")
def load_prefetched():
    """GitHub가 매일 새벽 받아둔 자료(cache 폴더)를 불러옴 → 서버 조회가 거의 없어져 빨라짐"""
    out = []
    for mod, f in ((dart, "dart.json.gz"), (prices, "prices.json.gz")):
        try:
            out.append(mod.load_cache(CACHE / f))
        except Exception:  # 파일이 없거나 코드가 옛 버전이면 그냥 실시간 조회
            out.append(0)
    return tuple(out)


load_prefetched()


@st.cache_data(ttl=86400, show_spinner="회사 목록 불러오는 중...")
def corp_list(key):
    f = CACHE / "corps.csv"
    if f.exists() and (dt.datetime.now().timestamp() - f.stat().st_mtime) < 7 * 86400:
        try:
            return pd.read_csv(f, dtype=str)
        except Exception:
            pass
    return dart.load_corp_codes(key)


@st.cache_data(ttl=21600, show_spinner="재무제표 불러오는 중...")
def fin(key, corp_code, years, q=0, fs="연결"):
    return dart.financials(key, corp_code, years, q, fs)


@st.cache_data(ttl=1800, show_spinner="공시 불러오는 중...")
def disc(key, corp_code, bgn, end):
    return dart.disclosures(key, corp_code, bgn, end)


@st.cache_data(ttl=3600, show_spinner="주가 불러오는 중...")
def price_hist(stock_code, days, gov_key):
    return prices.history(stock_code, days, gov_key)


@st.cache_data(ttl=21600)
def share_div(key, corp_code, year):
    s = d = None
    for y in (year, year - 1):  # 최신 사업보고서가 아직 없으면 전년
        try:
            s = s or dart.shares(key, corp_code, y)
            d = d if d is not None else dart.dividend(key, corp_code, y)
        except Exception:
            pass
    return s, d


@st.cache_data(ttl=21600, show_spinner="주가 불러오는 중...")
@st.cache_data(ttl=21600)
def disclosed_eps(key, corp_code, y):
    """y년 사업보고서 공시 기본EPS (없으면 한 해 전). (EPS, 사용 연도)"""
    for yy in (y, y - 1):
        try:
            e = dart.basic_eps(key, corp_code, yy)
        except Exception:
            e = None
        if e is not None:
            return e, yy
    return None, None


def valuation_for(key, corp_code, stock_code, year, q, end, row, gov_key, krx_key=""):
    price, mcap, lst, src = prices.price_on(stock_code, end, gov_key, krx_key)
    if price is None:
        return {}, None
    shares = lst
    if not shares:
        for y, rc in ((year, dart.Q_REPORT[q or 4]), (year, "11011"), (year - 1, "11011")):
            try:
                s = dart.shares(key, corp_code, y, rc)
            except Exception:
                s = None
            if s and s.get("발행주식수"):
                shares = s["발행주식수"]
                break
    try:
        dps = dart.dividend(key, corp_code, year if q in (0, 4) else year - 1)
    except Exception:
        dps = None
    f = 1 if q == 0 else 4  # 분기 실적은 연환산
    adj = dict(row)
    for k in ("지배주주순이익", "당기순이익", "매출액", "EBITDA"):
        if adj.get(k) is not None and not pd.isna(adj.get(k)):
            adj[k] = adj[k] * f
    # 가장 최근 확정 사업보고서 EPS: 연간(4Q 포함)은 그 해, 1~3분기는 전년도
    eps, _ = disclosed_eps(key, corp_code, year if q in (0, 4) else year - 1)
    v = prices.valuation(price, shares, adj, dps, mcap, eps)
    v.pop("주당배당금(원)", None)
    return v, src


def load_sectors():
    try:
        return json.loads((ROOT / "sectors.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def to_eok(df):
    out = df.copy()
    for c in set(dart.AMOUNTS) | set(dart.ITEMS):
        if c in out:
            out[c] = (pd.to_numeric(out[c], errors="coerce") / EOK).round(0)
    return out


def label(m):
    return f"{m} (억원)" if m in dart.AMOUNT_METRICS else m


def fmt(v):
    if v is None or pd.isna(v):
        return "-"
    return f"{v:,.2f}" if abs(v) < 1000 else f"{v:,.0f}"


MARKET = {"Y": "KOSPI", "K": "KOSDAQ", "N": "KONEX"}


@st.cache_data(ttl=86400)
def market_of(key, corp_code):
    """상장시장 (DART 기업개황 corp_cls)"""
    try:
        return MARKET.get(dart.company(key, corp_code).get("corp_cls"), "")
    except Exception:
        return ""


def html_table(rows, companies, markets=None):
    """지표명에 ? 윗첨자 + 마우스 올리면 계산 방법 표시"""
    esc = lambda x: html.escape(str(x), quote=True).replace("\n", "&#10;")
    markets = markets or {}
    head = "".join(f"<th>{esc(c)}" + (f"<br><span style='font-weight:400;font-size:12px;opacity:.7'>({markets[c]})</span>"
                                      if markets.get(c) else "") + "</th>" for c in companies)
    body, prev_g = [], None
    span = {}  # 구분별 줄 수 (연속된 같은 구분은 셀 합치기)
    for r in rows:
        span[r["구분"]] = span.get(r["구분"], 0) + 1
    for r in rows:
        k = r["_key"]
        tip = dart.DESC.get(k, "")
        name = esc(r["지표"]) + (f'<span class="tip" tabindex="0" data-tip="{esc(tip)}">?</span>' if tip else "")
        gcell = (f"<td class='g' rowspan='{span[r['구분']]}'>{esc(r['구분'])}</td>" if r["구분"] != prev_g else "")
        prev_g = r["구분"]
        f = (lambda v: "-" if v is None or pd.isna(v) else f"{v:,.0f}") if (k in dart.AMOUNT_METRICS or k in BASE_ROWS) else fmt
        cells = "".join(f"<td class='n'>{f(r[c])}</td>" for c in companies)
        body.append(f"<tr>{gcell}<td class='k'>{name}</td>{cells}</tr>")
    css = """<style>
.mt{border-collapse:collapse;width:100%;font-size:14px;table-layout:fixed}
.mt th,.mt td{border-bottom:1px solid rgba(128,128,128,.25);padding:6px 10px}
.mt td{white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.mt th{text-align:center;vertical-align:middle;font-weight:600;white-space:normal;word-break:keep-all;overflow-wrap:anywhere}
.mt td.n{text-align:right;font-variant-numeric:tabular-nums}.mt td.g{font-weight:600;opacity:.75;vertical-align:middle}
.mt td.k{overflow:visible}
.mt .tip{position:relative;display:inline-block;cursor:help;margin-left:4px;vertical-align:super;
 width:15px;height:15px;line-height:15px;text-align:center;border-radius:50%;font-size:10px;font-weight:700;
 background:rgba(128,128,128,.25);outline:none}
.mt .tip:hover,.mt .tip:focus{background:#ff4b4b;color:#fff}
.mt .tip:hover::after,.mt .tip:focus::after{content:attr(data-tip);position:absolute;left:22px;top:50%;
 transform:translateY(-50%);z-index:1000;white-space:pre-line;width:max-content;max-width:360px;
 padding:8px 10px;border-radius:6px;background:#262730;color:#fff;font-size:12.5px;font-weight:400;
 line-height:1.5;text-align:left;box-shadow:0 4px 12px rgba(0,0,0,.25);vertical-align:baseline}
</style>"""
    cols = "<col>" * len(companies)
    return (css + f"<div style='overflow-x:auto;padding-bottom:40px'><table class='mt' style='min-width:{280 + 90 * len(companies)}px'><colgroup><col style='width:84px'><col style='width:210px'>{cols}</colgroup><thead><tr><th>구분</th><th>지표</th>{head}</tr></thead>"
            f"<tbody>{''.join(body)}</tbody></table></div>")


def safe_msg(e):
    """오류 문구에서 인증키·주소를 지우고 짧게"""
    import re as _re
    s = _re.sub(r"(crtfc_key|serviceKey|AUTH_KEY)=[^&\s']+", r"\1=***", str(e))
    if "timed out" in s or "Max retries" in s or "Connection" in s:
        return "서버에 연결할 수 없어요 (시간 초과). 잠시 후 다시 시도해 주세요."
    return s[:200]


def _safe(f, *a):
    try:
        return f(*a)
    except Exception:
        return None


@st.cache_data(ttl=3600, show_spinner=False)
def excel_cached(sheets, footer):
    """같은 화면이면 엑셀 파일을 다시 만들지 않음"""
    return excel_bytes(sheets, footer)


def excel_bytes(sheets, footer=None):
    """엑셀 양식: 1행·A열 공란 · D3 틀고정 · 눈금선 없음 · 숫자 쉼표 · 음수 빨간색 · 헤더 가운데+연회색 · 원자료 C열 회사명"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    INT_FMT = "#,##0;[Red]-#,##0"
    DEC_FMT = "#,##0.00;[Red]-#,##0.00"
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            if df is None or df.empty:
                continue
            if name == "원자료" and "회사" in df:  # 회사명을 C열(세 번째)로
                q_ = ["분기"] if "분기" in df else []
                cols = [c for c in df.columns if c not in ["회사"] + q_]
                df = df[cols[:1] + ["회사"] + q_ + cols[1:]]  # A열 공란 → B=연도, C=회사, D=분기
            has_idx = not isinstance(df.index, pd.RangeIndex)
            df.to_excel(w, sheet_name=name[:31], index=has_idx, startrow=1, startcol=1)  # 1행·A열 비움
            ws = w.sheets[name[:31]]
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = "D3"
            hdr_rows = df.columns.nlevels
            for row in ws.iter_rows(min_row=2, max_row=1 + hdr_rows, min_col=2):  # A열(A2 등)은 색 없음
                for c in row:
                    c.fill = head_fill
                    c.alignment = center
            for row in ws.iter_rows(min_row=hdr_rows + 2):
                for c in row:
                    v = c.value
                    if isinstance(v, bool) or not isinstance(v, (int, float)):
                        continue
                    if name == "원자료" and c.column == 2:  # 연도는 쉼표 없이
                        c.number_format = "0"
                    else:
                        c.number_format = INT_FMT if float(v).is_integer() else DEC_FMT
            for row in ws.iter_rows():  # 글씨 크기 10
                for c in row:
                    c.font = Font(size=10, bold=c.font.bold)
            for i, col in enumerate(ws.columns, 1):  # 열 너비 대략 맞춤
                n = max((len(str(c.value)) for c in col if c.value is not None), default=4)
                ws.column_dimensions[get_column_letter(i)].width = 2 if i == 1 else min(max(n * 1.3 + 2, 8), 45)
        _link_formulas(w, INT_FMT, DEC_FMT)
        if footer and "지표비교" in w.sheets:  # 표 아래 한 줄 띄우고 기준 연도·분기 표시
            ws = w.sheets["지표비교"]
            c = ws.cell(ws.max_row + 2, 2, footer)
            c.font = Font(size=10)
    return buf.getvalue()


def _link_formulas(w, int_fmt, dec_fmt):
    """'지표비교' 값을 '원자료' 시트를 찾아오는 공식으로 바꿈 (회사명 + 지표 열 기준 INDEX/MATCH)"""
    from openpyxl.utils import get_column_letter
    if "지표비교" not in w.sheets or "원자료" not in w.sheets:
        return
    raw, cmp_ = w.sheets["원자료"], w.sheets["지표비교"]
    raw_col = {c.value: get_column_letter(c.column) for c in raw[2] if c.value is not None}  # 2행 = 헤더
    lab2key = {f"{k} (억원)": k for k in BASE_ROWS}
    lab2key.update({label(k): k for ks in dart.GROUPS.values() for k in ks})
    comp_col = raw_col.get("회사", "C")
    for r in range(3, cmp_.max_row + 1):
        k = lab2key.get(cmp_.cell(r, 3).value)  # C열 = 지표
        if k not in raw_col:
            continue
        for c in range(4, cmp_.max_column + 1):  # D열부터 회사
            cell = cmp_.cell(r, c)
            v = cell.value
            col = get_column_letter(c)
            look = f"INDEX('원자료'!${raw_col[k]}:${raw_col[k]},MATCH({col}$2,'원자료'!${comp_col}:${comp_col},0))"
            cell.value = f'=IFERROR(IF({look}="","",{look}),"")'
            is_int = isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer()
            cell.number_format = int_fmt if is_int else dec_fmt


# ================= 키 확인 =================
key = secret("DART_API_KEY")
gov_key = secret("DATA_GO_KR_KEY")
krx_key = secret("KRX_API_KEY")
if not key:
    key = st.text_input("DART 인증키", type="password")
    if not key:
        st.stop()
try:
    corps = corp_list(key)
except Exception as e:
    st.error(f"회사 목록을 못 불러왔어요: {safe_msg(e)}")
    st.stop()

name2row = {r.corp_name: r for r in corps.itertuples()}
all_labels = (corps["corp_name"] + " (" + corps["stock_code"] + ")").tolist()
lab2code = dict(zip(all_labels, corps["corp_code"]))
sectors = load_sectors()
code2row = {r.stock_code: r for r in corps.itertuples()}


def _norm(x):
    return x.replace(" ", "").replace("(주)", "").replace("㈜", "").lower()


def resolve(n):
    """sectors.json 항목(회사명 또는 종목코드) → 회사 목록. 정확히 → 종목코드 → 이름 일부 일치(여러 개면 모두)"""
    if n in name2row:
        return [name2row[n]]
    if n in code2row:
        return [code2row[n]]
    return [r for nm, r in name2row.items() if _norm(n) in _norm(nm)]

# ================= 헤더 =================
st.markdown("### 상장회사 분석")
DEFAULT_SECTOR = "패션" if "패션" in sectors else ALL
sector = st.pills("산업", [ALL] + list(sectors), default=DEFAULT_SECTOR, key="sector") or ALL

if sector == ALL:
    options, default = all_labels, []
else:
    resolved = {n: resolve(n) for n in sectors[sector]}
    found = list({r.stock_code: r for rs in resolved.values() for r in rs}.values())
    missing = [n for n, rs in resolved.items() if not rs]
    options = [f"{r.corp_name} ({r.stock_code})" for r in found]
    options += [l for l in all_labels if l not in options]  # 섹터 외 회사도 추가 가능
    default = options[:len(found)]
    if missing:
        st.caption(f"목록에서 못 찾은 회사(sectors.json 확인): {', '.join(missing)}")

h1, h2, h4 = st.columns([5, 2.4, 1.3], vertical_alignment="bottom")
picked = h1.multiselect("회사 (여러 개 선택 가능)", options, default=default, key=f"pick_{sector}_{abs(hash(tuple(default))) % 10**8}")  # 목록이 바뀌면 새로 선택
this_year = dt.date.today().year
y1, y2 = h2.slider("연도", 2015, this_year, (this_year - 3, this_year))
dl_slot = h4.empty()
years = tuple(range(y1, y2 + 1))

# ================= 왼쪽: 지표 선택 =================
st.sidebar.markdown("#### 표시할 재무지표")
DEFAULT_METRICS = ["영업이익률(%)", "부채비율(%)", "재고자산회전율(회)", "매출증가율(%)", "영업활동현금흐름", "PER(배)", "시가총액(억원)"]
BASE_ROWS = ["매출액", "영업이익"]  # 표 맨 위에 항상 표시 (억원)
chosen = []
for g, ks in dart.GROUPS.items():
    sel = st.sidebar.pills(g, ks, selection_mode="multi", default=[k for k in ks if k in DEFAULT_METRICS],
                           key=f"m_{g}",
                           format_func=label)
    chosen += sel or []
st.sidebar.caption("금액 단위: 억원 · 재무: 금융감독원 OpenDART")

if not picked:
    st.info("회사를 선택해 주세요.")
    st.stop()

tabs = st.tabs(["재무지표 비교", "주가·가치평가", "공시", "정기 리포트"])
QUARTERS = {"연간": 0, "1Q": 1, "2Q": 2, "3Q": 3, "4Q": 4}
with tabs[0]:
    c1, c2, c3 = st.columns([len(years), 5, 2.2], gap="medium")
    # 기본값: 가장 최근에 공시된 사업보고서 연도 (사업보고서는 보통 3월 말까지 제출)
    latest_ar = dart.eps_year_for(dt.date.today())
    by_default = latest_ar if latest_ar in years else years[-1]
    base_year = c1.segmented_control("기준 연도", list(years), default=by_default, key=f"by_{years}",
                                     format_func=lambda y: f"{y}년") or by_default
    qname = c2.segmented_control("분기", list(QUARTERS), default="연간", key="qtr") or "연간"
    fs = c3.segmented_control("재무제표", ["연결", "별도"], default="별도", key="fs",
                              help="연결: 자회사 포함 그룹 전체 · 별도: 회사 단독\n연결재무제표가 없는 회사는 별도로 표시") or "별도"
q = QUARTERS[qname]

# ================= 데이터 =================
# 여러 회사 자료를 동시에 미리 불러오기 (속도 향상) — 결과는 dart 모듈 안에 저장돼 아래에서 바로 사용
def _warm(code):
    for job in (lambda: dart.financials(key, code, (base_year,), q, fs),
                lambda: dart.company(key, code),
                lambda: dart.basic_eps(key, code, base_year if q in (0, 4) else base_year - 1),
                lambda: dart.dividend(key, code, base_year if q in (0, 4) else base_year - 1),
                lambda: dart.shares(key, code, base_year)):
        try:
            job()
        except Exception:
            pass


with ThreadPoolExecutor(max_workers=8) as ex:
    list(ex.map(_warm, [lab2code[l] for l in picked]))

data = {}
for lab in picked:
    try:
        data[lab] = fin(key, lab2code[lab], (base_year,), q, fs)
    except Exception as e:
        st.error(f"{lab}: {safe_msg(e)}")
        data[lab] = pd.DataFrame()

# 가치평가 지표 (기준 기간 말 종가 기준)
VAL_KEYS = dart.GROUPS["가치평가"]
PERIOD_END = {0: (12, 31), 1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}
val_end = min(dt.date(base_year, *PERIOD_END[q]), dt.date.today())
val_src = set()
if any(k in chosen for k in VAL_KEYS):
    for lab, d in data.items():
        if d.empty:
            continue
        stock = lab.split("(")[-1].rstrip(")")
        v, src = valuation_for(key, lab2code[lab], stock, base_year, q, val_end, d.iloc[-1].to_dict(), gov_key, krx_key)
        if src:
            val_src.add(src)
        for k, x in v.items():
            d.loc[d.index[-1], k] = x

frames = [to_eok(d).assign(회사=l.split(" (")[0]) for l, d in data.items() if not d.empty]
long = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


# ================= 1. 재무지표 비교 =================
table = pd.DataFrame()
with tabs[0]:
    if long.empty:
        st.warning(f"{base_year}년 {qname} {fs}재무제표 데이터가 없어요. (보고서가 아직 안 나왔거나 제출 대상이 아님)")
    elif not chosen:
        st.info("왼쪽에서 볼 지표를 골라 주세요.")
    else:
        snap = long[long["연도"] == base_year].set_index("회사")
        rows = [{"구분": "실적", "지표": f"{k} (억원)", "_key": k,
                 **{c: (snap.at[c, k] if k in snap else None) for c in snap.index}} for k in BASE_ROWS]
        for g, ks in dart.GROUPS.items():
            for k in ks:
                if k in chosen:
                    rows.append({"구분": g, "지표": label(k), "_key": k,
                                 **{c: (snap.at[c, k] if k in snap else None) for c in snap.index}})
        table = pd.DataFrame(rows).drop(columns="_key").set_index(["구분", "지표"])
        note = {0: "사업보고서", 4: "4Q = 사업보고서 연간 − 3분기 누적"}.get(q, f"{qname} 단일 분기")
        st.caption(f"{base_year}년 {qname} 기준 ({note}) · {fs}재무제표"
                   + (" · 분기 매출채권회전율·ROE·ROA는 연환산(×4), 증가율은 전년 동기 대비" if q else "")
                   + (f" · 가치평가는 {val_end:%Y-%m-%d} 종가 기준({'·'.join(sorted(val_src))})" if val_src else "")
                   + " · 지표의 ? 에 마우스를 올리면 계산 방법이 보여요")
        if q and any(k in chosen for k in ("재고자산회전율(회)", "재고자산회전일수(일)")):
            st.info("재고자산회전율(분기) 환산 방법: 최근 4개 분기 매출원가 합계 ÷ 평균 재고자산 "
                    "(1년 전 같은 분기말 재고와 이번 분기말 재고의 평균). "
                    "최근 4개 분기 매출원가 = 올해 누적 + 전년 연간 − 전년 같은 분기 누적 (4Q는 올해 연간 값).", icon="ℹ️")
        mk = {l.split(" (")[0]: market_of(key, lab2code[l]) for l in data}
        st.markdown(html_table(rows, list(snap.index), mk), unsafe_allow_html=True)
        if fs == "연결" and "기준" in snap:
            solo = [c for c in snap.index if snap.at[c, "기준"] == "별도"]
            if solo:
                st.caption(f"연결재무제표가 없어 별도로 표시한 회사: {', '.join(solo)}")
        fin_names = [c for c in snap.index if pd.isna(snap.at[c, "매출액"])] if "매출액" in snap else []
        if fin_names:
            st.caption(f"금융회사 등 매출액 계정이 없는 회사({', '.join(fin_names)})는 일부 지표가 빈칸이에요.")


# ================= 2. 주가·가치평가 =================
val_rows = []
with tabs[1]:
    period = st.select_slider("주가 기간", [90, 180, 365, 730, 1095], value=365,
                              format_func=lambda d: f"{d // 365}년" if d >= 365 else f"{d}일")
    # 주가·재무 자료를 여러 회사 동시에 미리 불러오기 (속도 향상)
    def _warm2(lab):
        code, cc = lab.split("(")[-1].rstrip(")"), lab2code[lab]
        for job in (lambda: prices.history(code, period, gov_key),
                    lambda: dart.financials(key, cc, (this_year - 2, this_year - 1), 0),
                    lambda: dart.shares(key, cc, this_year - 1),
                    lambda: dart.dividend(key, cc, this_year - 1),
                    lambda: dart.basic_eps(key, cc, dart.eps_year_for(dt.date.today()))):
            try:
                job()
            except Exception:
                pass
    with ThreadPoolExecutor(max_workers=8) as ex:
        list(ex.map(_warm2, picked))
    hist = []
    for lab in picked:
        code, name = lab.split("(")[-1].rstrip(")"), lab.split(" (")[0]
        try:
            ph, src = price_hist(code, period, gov_key)
        except Exception as e:
            st.error(f"{name} 주가: {safe_msg(e)}")
            continue
        if ph.empty:
            st.warning(f"{name}: 주가 데이터를 못 받았어요.")
            continue
        price = float(ph["종가"].iloc[-1])
        sh, dps = share_div(key, lab2code[lab], this_year - 1)
        mcap = float(ph["시가총액"].iloc[-1]) if "시가총액" in ph else None
        f_ann = fin(key, lab2code[lab], (this_year - 2, this_year - 1), 0)
        row = f_ann.iloc[-1] if not f_ann.empty else None
        eps, _ = disclosed_eps(key, lab2code[lab], dart.eps_year_for(ph["날짜"].iloc[-1].date()))
        if krx_key:  # 최신 종가는 KRX 공식 값으로
            k = prices.krx_price_on(code, ph["날짜"].iloc[-1].date(), krx_key)
            if k:
                price, mcap, src = k[0], k[1], "KRX"
        v = prices.valuation(price, (sh or {}).get("발행주식수"), row, dps, mcap, eps)
        val_rows.append({"회사": name, "기준일": f"{ph['날짜'].iloc[-1]:%Y-%m-%d}", "주가 출처": src, **v})
        hist.append(ph.assign(회사=name, 수익률=(ph["종가"] / ph["종가"].iloc[0] - 1) * 100))
    if val_rows:
        vt = pd.DataFrame(val_rows).set_index("회사")
        show = vt.T.map(lambda v: fmt(v) if isinstance(v, (int, float)) else ("-" if v is None else str(v)))
        st.dataframe(show, width="stretch")
        st.caption("PER = 종가 ÷ 직전 사업연도 공시 기본EPS (KRX 방식). PBR은 지배주주자본 ÷ 보통주 발행주식수(근사치). 적자면 PER 표시 안 함.")
    if hist:
        h = pd.concat(hist)
        mode = st.segmented_control("주가 차트", ["수익률(%)", "주가(원)"], default="수익률(%)") or "수익률(%)"
        y = "수익률" if mode.startswith("수익률") else "종가"
        st.plotly_chart(px.line(h, x="날짜", y=y, color="회사", title=f"주가 비교 · {mode}"), width="stretch")
val_df = pd.DataFrame(val_rows)

# ================= 3. 공시 =================
dis = pd.DataFrame()
with tabs[2]:
    days = st.select_slider("기간", [7, 30, 90, 180, 365], value=30, format_func=lambda d: f"최근 {d}일")
    only_imp = st.checkbox("중요 공시만 보기 (증자·전환사채·최대주주 변경·소송 등)")
    end = dt.date.today()
    bgn = end - dt.timedelta(days=days)
    with ThreadPoolExecutor(max_workers=8) as ex:  # 공시 동시 조회
        list(ex.map(lambda l: _safe(dart.disclosures, key, lab2code[l], bgn.strftime("%Y%m%d"), end.strftime("%Y%m%d")), picked))
    parts = []
    for lab in picked:
        try:
            parts.append(disc(key, lab2code[lab], bgn.strftime("%Y%m%d"), end.strftime("%Y%m%d")))
        except Exception as e:
            st.error(f"{lab}: {safe_msg(e)}")
    parts = [p for p in parts if not p.empty]
    if parts:
        dis = pd.concat(parts)
        dis.insert(0, "중요", dis["공시명"].map(lambda t: "●" if dart.is_important(t) else ""))
        if only_imp:
            dis = dis[dis["중요"] == "●"]
        dis = dis.sort_values("접수일", ascending=False)
        st.dataframe(dis, width="stretch", hide_index=True,
                     column_config={"링크": st.column_config.LinkColumn("원문", display_text="열기")})
    else:
        st.info("해당 기간 공시가 없어요.")

# ================= 4. 정기 리포트 =================
with tabs[3]:
    rdir = ROOT / "reports"
    files = sorted([f for f in rdir.glob("*.md") if f.name != "latest.md"], reverse=True) if rdir.exists() else []
    if not files:
        st.info("아직 리포트가 없어요. GitHub Actions가 매일 아침 자동으로 만들어요. (README 참고)")
    else:
        f = st.selectbox("날짜", files, format_func=lambda p: p.stem)
        st.markdown(f.read_text(encoding="utf-8"))

# ================= 엑셀 (헤더 버튼) =================
raw_x = long.assign(분기=qname) if not long.empty else long
sheets = {"지표비교": table, "원자료": raw_x, "가치평가": val_df, "공시": dis}
dl_slot.download_button("엑셀 받기", excel_cached(sheets, f"기준: {base_year}년 {qname} · {fs}재무제표"),
                        file_name=f"상장사분석_{base_year}년_{qname}_{dt.date.today():%Y%m%d}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        width="stretch")
